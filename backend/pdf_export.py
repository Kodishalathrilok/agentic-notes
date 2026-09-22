"""
Export helpers: render generated notes/quiz/flashcards to Markdown or PDF.
"""

import re
import csv
from io import BytesIO, StringIO

from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
)
from xml.sax.saxutils import escape


def notes_to_markdown(notes: str, quiz: str = "", flashcards: str = "") -> str:
    """Join the three sections into a single Markdown document."""
    parts = []
    if notes and notes.strip():
        parts.append("# AI Generated Notes\n\n" + notes.strip())
    if quiz and quiz.strip():
        parts.append("## Quiz\n\n" + quiz.strip())
    if flashcards and flashcards.strip():
        parts.append("## Flashcards\n\n" + flashcards.strip())
    return "\n\n---\n\n".join(parts) + "\n"


# ---------------------------------------------------------------------------
# Flashcards -> CSV (Anki / Quizlet compatible)
# ---------------------------------------------------------------------------

def parse_flashcards(raw: str):
    """Parse the CARD/Front/Back text format into [{front, back}, ...]."""
    cards = []
    current = None
    mode = None

    def push():
        if current and (current["front"] or current["back"]):
            cards.append(current)

    for line in (raw or "").splitlines():
        s = line.strip()
        if not s:
            continue
        if re.match(r"^CARD\s*\d+", s, re.IGNORECASE):
            push()
            current = {"front": "", "back": ""}
            mode = None
            continue
        fm = re.match(r"^Front\s*:?\s*(.*)$", s, re.IGNORECASE)
        bm = re.match(r"^Back\s*:?\s*(.*)$", s, re.IGNORECASE)
        if fm:
            if current is None:
                current = {"front": "", "back": ""}
            current["front"] = fm.group(1).strip()
            mode = "front"
        elif bm:
            if current is None:
                current = {"front": "", "back": ""}
            current["back"] = bm.group(1).strip()
            mode = "back"
        elif current and mode:
            current[mode] += (" " if current[mode] else "") + s
    push()
    return cards


# CSV/formula injection: a cell starting with one of these is evaluated as a
# formula by Excel, LibreOffice and Google Sheets (e.g. =HYPERLINK(...)).
_ALWAYS_FORMULA = ("=", "@", "\t", "\r")


def _csv_safe(cell: str) -> str:
    """Neutralise a spreadsheet formula by prefixing an apostrophe (OWASP).

    `=`, `@`, tab and CR always start a formula. `+` and `-` are escaped
    when NOT followed by a digit, `.` or a space (`-cmd|...`,
    `+HYPERLINK(...)`), so ordinary values such as `-38.8 °C`, `+3` or
    `- item` stay untouched. A number-like start is still escaped if the cell
    contains what turns arithmetic into code - a function call `(`, a DDE
    pipe `|` or a sheet reference `!` (e.g. `-2+3+cmd|' /C calc'!A0`).
    """
    if cell.startswith(_ALWAYS_FORMULA):
        return "'" + cell
    if cell[:1] in ("+", "-") and len(cell) > 1:
        if not (cell[1].isdigit() or cell[1] in ". ") or any(c in cell for c in "(|!"):
            return "'" + cell
    return cell


def flashcards_to_csv(flashcards: str) -> str:
    """Render flashcards as a two-column CSV (front,back) for Anki/Quizlet.

    Card text comes from a model that was fed user-supplied documents, so a
    cell could carry a formula; every cell goes through _csv_safe.
    """
    buf = StringIO()
    writer = csv.writer(buf)
    for card in parse_flashcards(flashcards):
        writer.writerow([_csv_safe(card["front"]), _csv_safe(card["back"])])
    return buf.getvalue()


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def _inline_bold(text: str) -> str:
    """Convert **bold** markers to reportlab <b> tags (after escaping)."""
    safe = escape(text)
    # **bold** -> <b>bold</b>
    return re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", safe)


def _add_block(story, raw_text, styles):
    """Render a block of notes/quiz/flashcard text into reportlab flowables."""
    body = styles["body"]
    header = styles["header"]

    for line in raw_text.splitlines():
        stripped = line.strip()
        if not stripped:
            story.append(Spacer(1, 6))
            continue

        # Bold-only header line: **Header:**  or  **Header**
        m = re.fullmatch(r"\*\*(.+?):?\*\*", stripped)
        if m:
            story.append(Spacer(1, 4))
            story.append(Paragraph(escape(m.group(1)), header))
            continue

        # Bullets
        if stripped.startswith("•"):
            content = _inline_bold(stripped[1:].strip())
            story.append(Paragraph(f"&bull;&nbsp;&nbsp;{content}", body))
            continue

        # Sub-bullets
        if stripped.startswith("-"):
            content = _inline_bold(stripped[1:].strip())
            story.append(Paragraph(f"&nbsp;&nbsp;&nbsp;&nbsp;&ndash;&nbsp;{content}", body))
            continue

        # Numbered
        if re.match(r"^\d+[\.\)]", stripped):
            story.append(Paragraph(_inline_bold(stripped), body))
            continue

        # Plain paragraph
        story.append(Paragraph(_inline_bold(stripped), body))


def notes_to_docx(notes: str, quiz: str = "", flashcards: str = "") -> bytes:
    """Render the sections to a .docx Word document and return the bytes."""
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()

    title = doc.add_heading("AI Generated Notes", level=0)
    for run in title.runs:
        run.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)

    def add_block(raw):
        for line in raw.splitlines():
            s = line.strip()
            if not s:
                continue
            header = re.fullmatch(r"\*\*(.+?):?\*\*", s)
            if header:
                doc.add_heading(header.group(1), level=2)
                continue
            text = re.sub(r"\*\*(.+?)\*\*", r"\1", s)  # drop bold markers
            if s.startswith("•") or s.startswith("-"):
                p = doc.add_paragraph(text.lstrip("•- ").strip(), style="List Bullet")
            elif re.match(r"^\d+[.)]", s):
                p = doc.add_paragraph(text, style="List Number")
            else:
                p = doc.add_paragraph(text)
            for run in p.runs:
                run.font.size = Pt(10)

    add_block(notes)

    if quiz and quiz.strip():
        doc.add_heading("Quiz", level=1)
        add_block(quiz)
    if flashcards and flashcards.strip():
        doc.add_heading("Flashcards", level=1)
        add_block(flashcards)

    buffer = BytesIO()
    doc.save(buffer)
    data = buffer.getvalue()
    buffer.close()
    return data


def notes_to_pdf(notes: str, quiz: str = "", flashcards: str = "") -> bytes:
    """Render the sections to a PDF and return the raw bytes."""
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=2 * cm,
        rightMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
        title="AI Generated Notes",
    )

    base = getSampleStyleSheet()
    styles = {
        "title": ParagraphStyle(
            "TitleBlue",
            parent=base["Title"],
            textColor=HexColor("#1e40af"),
            fontSize=22,
            leading=26,
            spaceAfter=14,
        ),
        "section": ParagraphStyle(
            "SectionHeader",
            parent=base["Heading2"],
            textColor=HexColor("#374151"),
            fontSize=15,
            leading=20,
            spaceBefore=12,
            spaceAfter=6,
        ),
        "header": ParagraphStyle(
            "InlineHeader",
            parent=base["Heading3"],
            textColor=HexColor("#374151"),
            fontSize=12,
            leading=16,
            spaceBefore=6,
            spaceAfter=2,
        ),
        "body": ParagraphStyle(
            "BodyTextCustom",
            parent=base["BodyText"],
            fontSize=10,
            leading=16,
            alignment=TA_LEFT,
            textColor=HexColor("#111827"),
        ),
    }

    story = [Paragraph("AI Generated Notes", styles["title"])]

    if notes and notes.strip():
        _add_block(story, notes, styles)

    if quiz and quiz.strip():
        story.append(Spacer(1, 10))
        story.append(Paragraph("Quiz", styles["section"]))
        _add_block(story, quiz, styles)

    if flashcards and flashcards.strip():
        story.append(Spacer(1, 10))
        story.append(Paragraph("Flashcards", styles["section"]))
        _add_block(story, flashcards, styles)

    doc.build(story)
    pdf_bytes = buffer.getvalue()
    buffer.close()
    return pdf_bytes
